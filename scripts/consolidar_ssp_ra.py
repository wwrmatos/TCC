# /// script
# requires-python = ">=3.13"
# dependencies = ["polars>=1.30", "openpyxl>=3.1", "pdfplumber>=0.11"]
# ///
"""Consolida os balanços criminais do DF e das RAs baixados por
`baixar_ssp_ra.py` (`data/ssp_df_ra_bruto/`) num único CSV com todos os anos:
`data/processed/ssp_df_ra_consolidado.csv`.

Uma linha por RA, ano, natureza e contagem, no mesmo formato de `ssp_dp.csv`
(`ra, ano, eixo, natureza, total, jan … dez`), mais `ra_numero` (0 = DF
inteiro), `contagem` (OCORRÊNCIA ou VÍTIMA — só o DF traz vítimas, e só nas
naturezas de C.V.L.I.) e as colunas de procedência `fonte` (xlsx/pdf) e `arquivo`.

Para cada RA/ano usa a planilha quando existe; senão, o PDF (em 2018, o "PDF2",
que tem o mesmo layout dos anos seguintes). Antes de usar um arquivo, confere o
número da RA no cabeçalho: o site tem arquivos trocados (ex.: o de Sobradinho I
2015 é o de Sobradinho II), e esses são ignorados.

Nos PDFs, a tabela é lida pelas bordas desenhadas (linhas e coluna NATUREZA) e
pela posição dos números sob o cabeçalho de cada mês; célula vazia vale 0. O eixo
das naturezas vem das planilhas, já que no PDF ele é texto girado. Meses a partir
da data de atualização indicada na nota "Obs" ficam nulos (o ano corrente vem
com 0 nos meses futuros).

Script independente do projeto (dependências declaradas acima). Uso:
    uv run --script scripts/consolidar_ssp_ra.py
"""

import re
import unicodedata
from pathlib import Path

import openpyxl
import pdfplumber
import polars as pl

RAIZ = Path(__file__).resolve().parents[1]
BRUTO_DIR = RAIZ / "data" / "ssp_df_ra_bruto"
SAIDA_CSV = RAIZ / "data" / "processed" / "ssp_df_ra_consolidado.csv"

MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

# Número oficial de cada RA e o nome usado na saída (0 = o DF inteiro).
RAS: dict[int, str] = {
    0: "DISTRITO FEDERAL",
    1: "BRASÍLIA", 2: "GAMA", 3: "TAGUATINGA", 4: "BRAZLÂNDIA", 5: "SOBRADINHO",
    6: "PLANALTINA", 7: "PARANOÁ", 8: "NÚCLEO BANDEIRANTE", 9: "CEILÂNDIA", 10: "GUARÁ",
    11: "CRUZEIRO", 12: "SAMAMBAIA", 13: "SANTA MARIA", 14: "SÃO SEBASTIÃO",
    15: "RECANTO DAS EMAS", 16: "LAGO SUL", 17: "RIACHO FUNDO", 18: "LAGO NORTE",
    19: "CANDANGOLÂNDIA", 20: "ÁGUAS CLARAS", 21: "RIACHO FUNDO II", 22: "SUDOESTE/OCTOGONAL",
    23: "VARJÃO", 24: "PARK WAY", 25: "SCIA/ESTRUTURAL", 26: "SOBRADINHO II",
    27: "JARDIM BOTÂNICO", 28: "ITAPOÃ", 29: "SIA", 30: "VICENTE PIRES", 31: "FERCAL",
    32: "SOL NASCENTE/PÔR DO SOL", 33: "ARNIQUEIRA",
}

# Nome da RA na página da SSP (coluna `ra` do manifesto) → número da RA.
RA_DO_SITE: dict[str, int] = {
    "DISTRITO FEDERAL": 0,
    "BRASÍLIA (PLANO PILOTO)": 1, "GAMA": 2, "TAGUATINGA": 3, "BRAZLÂNDIA": 4,
    "SOBRADINHO I": 5, "PLANALTINA": 6, "PARANOÁ": 7, "NÚCLEO BANDEIRANTE": 8,
    "CEILÂNDIA": 9, "GUARÁ": 10, "CRUZEIRO": 11, "SAMAMBAIA": 12, "SANTA MARIA": 13,
    "SÃO SEBASTIÃO": 14, "RECANTO DAS EMAS": 15, "LAGO SUL": 16, "RIACHO FUNDO I": 17,
    "LAGO NORTE": 18, "CANDANGOLÂNDIA": 19, "ÁGUAS CLARAS": 20, "RIACHO FUNDO II": 21,
    "SUDOESTE/OCTOGONAL": 22, "VARJÃO": 23, "PARK WAY": 24,
    "ESTRUTURAL / SETOR COMPLEMENTAR DE INDÚSTRIA E ABASTECIMENTO (SCIA)": 25,
    "SOBRADINHO II": 26, "JARDIM BOTÂNICO": 27, "ITAPOÃ": 28, "SIA": 29,
    "VICENTE PIRES": 30, "FERCAL": 31, "SOL NASCENTE": 32, "ARNIQUEIRA": 33,
}

_ROMANOS = {"I": 1, "V": 5, "X": 10, "L": 50}


def _romano(s: str) -> int:
    valores = [_ROMANOS[c] for c in s]
    return sum(-v if i + 1 < len(valores) and v < valores[i + 1] else v for i, v in enumerate(valores))


def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _limpar(s: str) -> str:
    return " ".join(str(s).split()).upper()


def ra_do_cabecalho(texto: str) -> int | None:
    """Número da RA citado no cabeçalho ("RA XXV - SCIA/ESTRUTURAL")."""
    if m := re.search(r"\bRA ([IVXL]+) ?-", texto):
        return _romano(m.group(1))
    # Arniqueira e Sol Nascente aparecem às vezes sem o número.
    sem = _sem_acento(texto.upper())
    if "ARNIQUEIRA" in sem:
        return 33
    if "SOL NASCENTE" in sem:
        return 32
    # O balanço do DF: "DISTRITO FEDERAL" sozinho numa linha (todo arquivo tem
    # "GOVERNO DO DISTRITO FEDERAL" no topo, por isso a linha inteira).
    if re.search(r"^(PRINCIPAIS OCORRENCIAS POLICIAIS NO )?DISTRITO FEDERAL$", sem, re.M):
        return 0
    return None


def meses_validos(texto: str, ano: int) -> int:
    """Quantos meses do ano têm dado fechado, pela data de atualização da nota "Obs"."""
    if m := re.search(r"atualizados em \d{1,2}/(\d{1,2})/(\d{2,4})", texto):
        mes, ano_atu = int(m.group(1)), int(m.group(2))
        ano_atu += 2000 if ano_atu < 100 else 0
        if ano_atu == ano:
            return mes - 1
    return 12


# --------------------------------------------------------------------------- xlsx

def ler_xlsx(path: Path) -> tuple[str, list[dict]]:
    """Lê a planilha de uma RA ou do DF.

    A do DF tem uma coluna a mais, entre eixo e natureza, dizendo se a linha
    conta OCORRÊNCIA ou VÍTIMA; a linha de vítimas vem logo abaixo da de
    ocorrências, com a natureza em branco."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    linhas = list(ws.iter_rows(values_only=True))
    texto = "\n".join(str(c) for row in linhas for c in row if isinstance(c, str))

    i_cab = next(i for i, row in enumerate(linhas) if "NATUREZA" in [_limpar(c or "") for c in row])
    cab = [_limpar(c or "") for c in linhas[i_cab]]
    col_nat, col_total = cab.index("NATUREZA"), cab.index("TOTAL")
    col_tipo = 1 if col_nat == 2 else None
    i_meses = i_cab + 1
    col_mes = {
        _limpar(c): j for j, c in enumerate(linhas[i_meses]) if isinstance(c, str) and _limpar(c) in MESES
    }

    registros, eixo, natureza = [], None, None
    for row in linhas[i_meses + 1:]:
        if row[0] is not None and str(row[0]).strip():
            eixo = _limpar(row[0])
        if "TOTAL" in (eixo or "") or _limpar(row[0] or "").startswith(("FONTE", "OBS", "*")):
            continue  # linhas de total do eixo e notas de rodapé
        tipo = _limpar(row[col_tipo] or "") if col_tipo is not None else ""
        if row[col_nat] is not None and str(row[col_nat]).strip():
            natureza = _limpar(row[col_nat])
        elif tipo != "VÍTIMA":
            continue  # linha vazia
        registros.append({
            "eixo": eixo,
            "natureza": natureza,
            "contagem": "VÍTIMA" if tipo == "VÍTIMA" else "OCORRÊNCIA",
            "total": row[col_total],
            **{m.lower(): row[j] for m, j in col_mes.items()},
        })
    return texto, registros


# ---------------------------------------------------------------------------- pdf

def _agrupar(valores: list[float], tol: float = 2.5) -> list[float]:
    grupos: list[float] = []
    for v in sorted(valores):
        if not grupos or v - grupos[-1] > tol:
            grupos.append(v)
    return grupos


def ler_pdf(path: Path) -> tuple[str, list[dict]]:
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[0]
        texto = page.extract_text() or ""
        words = page.extract_words()

        # Cabeçalho: "NATUREZA" e "TOTAL" na mesma linha (o título também diz
        # "POR NATUREZA", daí exigir o TOTAL ao lado).
        nat, tot = next(
            (n, t) for n in words if n["text"] == "NATUREZA"
            for t in words if t["text"] == "TOTAL" and abs(t["top"] - n["top"]) < 10 and t["x0"] > n["x0"]
        )
        cab_meses = {}
        for w in words:
            if w["text"] in MESES and nat["top"] - 5 < w["top"] < nat["top"] + 30:
                cab_meses.setdefault(w["text"], w)
        centros = {"total": (tot["x0"] + tot["x1"]) / 2}
        centros |= {m.lower(): (w["x0"] + w["x1"]) / 2 for m, w in cab_meses.items()}
        fim_cab = max(w["bottom"] for w in cab_meses.values())

        # Coluna NATUREZA: entre a borda vertical à esquerda do título "NATUREZA"
        # e a borda à esquerda de "TOTAL".
        xs = _agrupar([e["x0"] for e in page.vertical_edges])
        nat_esq = max(x for x in xs if x <= nat["x0"])
        nat_dir = max(x for x in xs if x <= tot["x0"])

        # Linhas da tabela: bordas horizontais que atravessam a coluna TOTAL.
        ys = _agrupar([
            e["top"] for e in page.horizontal_edges
            if e["x0"] <= centros["total"] <= e["x1"] and e["top"] > fim_cab - 2
        ])

        def tem_borda_vertical(x: float, y: float) -> bool:
            return any(
                abs(e["x0"] - x) < 2.5 and e["top"] - 1 <= y <= e["bottom"] + 1
                for e in page.vertical_edges
            )

        registros = []
        for y0, y1 in zip(ys, ys[1:]):
            meio = (y0 + y1) / 2
            # Linhas de total ("1.TOTAL C.V.L.I.") mesclam eixo e natureza: sem
            # borda à esquerda da coluna NATUREZA.
            if not tem_borda_vertical(nat_esq, meio):
                continue
            natureza = _limpar(page.crop((nat_esq, y0, nat_dir, y1)).extract_text() or "")
            if not natureza or natureza.startswith("TOTAL"):
                continue
            # Alguns PDFs espaçam os dígitos ("1 7 5 7"): junta, em ordem, os
            # pedaços numéricos que caem na mesma coluna.
            pedacos: dict[str, list[tuple[float, str]]] = {}
            for w in words:
                cy = (w["top"] + w["bottom"]) / 2
                if y0 < cy < y1 and w["x0"] >= nat_dir - 1 and re.fullmatch(r"[\d.]+", w["text"]):
                    cx = (w["x0"] + w["x1"]) / 2
                    col = min(centros, key=lambda c: abs(centros[c] - cx))
                    pedacos.setdefault(col, []).append((w["x0"], w["text"]))
            valores = {
                col: int("".join(t for _, t in sorted(ps)).replace(".", ""))
                for col, ps in pedacos.items()
            }
            if "total" not in valores:
                continue
            registros.append({
                "eixo": None,
                "natureza": natureza,
                "contagem": "OCORRÊNCIA",
                "total": valores["total"],
                **{m.lower(): valores.get(m.lower(), 0) for m in MESES},
            })
    return texto, registros


# --------------------------------------------------------------------- consolidação

def corrigir_rotulos(registros: list[dict]) -> list[dict]:
    """Em alguns arquivos (2016–2018) a linha de TENTATIVA DE LATROCÍNIO, logo
    após TENTATIVA DE HOMICÍDIO, veio rotulada como "ROUBO EM RESIDÊNCIA",
    repetindo a natureza do eixo C.C.P."""
    nomes = [r["natureza"] for r in registros]
    if "TENTATIVA DE LATROCÍNIO" not in nomes and nomes.count("ROUBO EM RESIDÊNCIA") > 1:
        for anterior, r in zip(registros, registros[1:]):
            if anterior["natureza"] == "TENTATIVA DE HOMICÍDIO" and r["natureza"] == "ROUBO EM RESIDÊNCIA":
                r["natureza"] = "TENTATIVA DE LATROCÍNIO"
    return registros


def ano_do_cabecalho(texto: str) -> int | None:
    m = re.search(r"COMPARATIVO MENSAL (\d{4})", texto)
    return int(m.group(1)) if m else None


def candidatos(arquivos: list[dict]) -> list[dict]:
    """Ordem de preferência: planilha, depois PDF2 (2018), depois os demais PDFs."""
    def chave(a: dict) -> tuple:
        return (a["tipo"] != "xlsx", not a["rotulo"].upper().replace(" ", "").startswith("PDF2"))
    return sorted(arquivos, key=chave)


# Mesma natureza com nomes diferentes entre o balanço do DF e o das RAs; fica o das RAs.
SINONIMOS = {
    "ROUBO EM TRANSPORTE COLETIVO": "ROUBO EM COLETIVO",
    "POSSE/PORTE DE ARMA DE FOGO": "POSSE/PORTE DE ARMA",
}


def canonizar_naturezas(df: pl.DataFrame) -> pl.DataFrame:
    """Unifica sinônimos e grafias que diferem só em acento (fica a mais frequente)."""
    # Asteriscos só remetem a notas de rodapé ("ROUBO EM COMÉRCIO *", "HOMICÍDIO *").
    df = df.with_columns(pl.col("natureza").str.replace(r"\s*\*+$", "").replace(SINONIMOS))
    freq = df.group_by("natureza").len()
    canonica: dict[str, tuple[str, int]] = {}
    for nat, n in freq.iter_rows():
        chave = _sem_acento(nat)
        if chave not in canonica or n > canonica[chave][1]:
            canonica[chave] = (nat, n)
    return df.with_columns(
        pl.col("natureza").map_elements(lambda s: canonica[_sem_acento(s)][0], return_dtype=pl.String)
    )


def main() -> None:
    # Mantém também os links descartados no download, para que a RA/ano sem
    # nenhum arquivo apareça no relatório de faltantes.
    manifesto = pl.read_csv(BRUTO_DIR / "manifesto.csv")

    tabelas, faltando, ignorados = [], [], []
    for (ra_site, ano), grupo in manifesto.group_by("ra", "ano", maintain_order=True):
        esperado = RA_DO_SITE[ra_site]
        usado = None
        for arq in candidatos(grupo.filter(pl.col("arquivo").is_not_null()).to_dicts()):
            path = BRUTO_DIR / arq["arquivo"]
            try:
                texto, registros = (ler_xlsx if arq["tipo"] == "xlsx" else ler_pdf)(path)
            except Exception as e:  # noqa: BLE001 — tenta o próximo arquivo
                ignorados.append((arq["arquivo"], f"erro de leitura: {e!r}"))
                continue
            ra_arq, ano_arq = ra_do_cabecalho(texto), ano_do_cabecalho(texto)
            if ra_arq != esperado or (ano_arq and ano_arq != ano):
                ignorados.append((arq["arquivo"], f"cabeçalho diz RA {ra_arq}, ano {ano_arq}"))
                continue
            if not registros:
                ignorados.append((arq["arquivo"], "nenhuma linha lida"))
                continue
            n_meses = meses_validos(texto, ano)
            df = pl.DataFrame(corrigir_rotulos(registros), infer_schema_length=None).with_columns(
                [pl.col(m.lower()).cast(pl.Int64, strict=False) for m in MESES]
                + [pl.col("total").cast(pl.Int64, strict=False)]
            ).with_columns(
                # Célula vazia num mês já fechado é zero; meses ainda não
                # fechados ficam nulos.
                [pl.col(m.lower()).fill_null(0) for m in MESES[:n_meses]]
                + [pl.lit(None, pl.Int64).alias(m.lower()) for m in MESES[n_meses:]]
            )
            tabelas.append(df.with_columns(
                ra_numero=pl.lit(esperado),
                ra=pl.lit(RAS[esperado]),
                ano=pl.lit(ano),
                fonte=pl.lit(arq["tipo"]),
                arquivo=pl.lit(arq["arquivo"]),
            ))
            usado = arq["arquivo"]
            break
        if usado is None:
            faltando.append((ano, ra_site))

    df = pl.concat(tabelas, how="diagonal_relaxed")
    df = canonizar_naturezas(df)

    # Eixo: vem das planilhas; os PDFs herdam pelo nome da natureza.
    eixo_por_natureza = dict(
        df.filter(pl.col("eixo").is_not_null())
        .group_by("natureza").agg(pl.col("eixo").mode().first())
        .iter_rows()
    )
    df = df.with_columns(
        pl.col("eixo").fill_null(pl.col("natureza").replace_strict(eixo_por_natureza, default=None))
    )

    colunas = ["ra_numero", "ra", "ano", "eixo", "natureza", "contagem", "total",
               *[m.lower() for m in MESES], "fonte", "arquivo"]
    df = df.select(colunas).sort("ano", "ra_numero")

    # Checagem: soma dos meses = total (só onde todos os meses estão preenchidos).
    meses = [m.lower() for m in MESES]
    soma_diverge = df.filter(
        pl.all_horizontal(pl.col(meses).is_not_null())
        & (pl.sum_horizontal(meses) != pl.col("total"))
    )

    SAIDA_CSV.parent.mkdir(parents=True, exist_ok=True)
    df.write_csv(SAIDA_CSV)

    print(f"{df.height} linhas em {SAIDA_CSV.relative_to(RAIZ)}")
    with pl.Config(tbl_rows=-1, fmt_str_lengths=70, tbl_width_chars=200):
        print(df.group_by("ano").agg(
            locais=pl.col("ra").n_unique(),  # RAs + DF
            de_xlsx=pl.col("ra").filter(pl.col("fonte") == "xlsx").n_unique(),
            de_pdf=pl.col("ra").filter(pl.col("fonte") == "pdf").n_unique(),
            naturezas=pl.col("natureza").n_unique(),
        ).sort("ano"))
        sem_eixo = df.filter(pl.col("eixo").is_null())["natureza"].unique().to_list()
        if sem_eixo:
            print("Naturezas sem eixo (só aparecem em PDF):", sem_eixo)
        if soma_diverge.height:
            # Vem da fonte (o TOTAL impresso não fecha com os meses), não da leitura.
            print(f"{soma_diverge.height} linhas em que a soma dos meses ≠ total (erro da fonte):")
            print(soma_diverge.group_by("ano", "fonte").agg(
                linhas=pl.len(), ras=pl.col("ra").unique().sort().str.join(", ")
            ).sort("ano"))
    if ignorados:
        print("Arquivos ignorados:")
        for arq, motivo in ignorados:
            print(f"  {arq}: {motivo}")
    if faltando:
        print("RA/ano sem nenhum arquivo utilizável:")
        for ano, ra in sorted(faltando):
            print(f"  {ano} {ra}")


if __name__ == "__main__":
    main()
